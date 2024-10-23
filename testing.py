from sqlite3 import Error
from unittest.mock import MagicMock, patch
from unittest.mock import MagicMock, patch
import pytest
from pptx import Presentation

from data_extractor.storage.sql_storage import SQLStorage




from unittest.mock import MagicMock
from data_extractor.data_extractor.extractor import FileExtractor  # Adjust the import according to your module

@pytest.fixture
def setup_extractor():
    mock_loader = MagicMock()
    extractor = FileExtractor(loader=mock_loader, file_type='pdf')
    return extractor, mock_loader

def test_extract_text_empty_pdf(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_loader.load_file.return_value = MagicMock(pages=[])
    extractor.load('empty.pdf')
    assert extractor.extract_text() == ''

def test_extract_text_single_page(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Hello, World!"
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page])
    extractor.load('single_page.pdf')
    assert extractor.extract_text() == "Hello, World!"

def test_extract_text_multiple_pages(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Page 1 text."
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Page 2 text."
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page1, mock_page2])
    extractor.load('multi_page.pdf')
    assert extractor.extract_text() == "Page 1 text.Page 2 text."

def test_extract_text_with_newlines(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Line 1.\nLine 2."
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page])
    extractor.load('newline_pdf.pdf')
    assert extractor.extract_text() == "Line 1.\nLine 2."

def test_extract_text_with_special_characters(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Café, résumé, and naïve."
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page])
    extractor.load('special_characters.pdf')
    assert extractor.extract_text() == "Café, résumé, and naïve."

def test_extract_text_pdf_with_no_text(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page = MagicMock()
    mock_page.extract_text.return_value = ""
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page])
    extractor.load('no_text_pdf.pdf')
    assert extractor.extract_text() == ""

def test_extract_text_pdf_with_empty_and_non_empty_pages(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_empty_page = MagicMock()
    mock_empty_page.extract_text.return_value = ""
    mock_non_empty_page = MagicMock()
    mock_non_empty_page.extract_text.return_value = "Some text."
    mock_loader.load_file.return_value = MagicMock(pages=[mock_empty_page, mock_non_empty_page])
    extractor.load('mixed_pages.pdf')
    assert extractor.extract_text() == "Some text."

def test_extract_text_large_pdf(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_pages = [MagicMock(extract_text=MagicMock(return_value=f"Page {i} text.")) for i in range(100)]
    mock_loader.load_file.return_value = MagicMock(pages=mock_pages)
    extractor.load('large_pdf.pdf')
    expected_text = ''.join(f"Page {i} text." for i in range(100))
    assert extractor.extract_text() == expected_text

def test_extract_text_pdf_with_tables(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Table data."
    mock_loader.load_file.return_value = MagicMock(pages=[mock_page])
    extractor.load('table_pdf.pdf')
    assert extractor.extract_text() == "Table data."

def test_extract_text_pdf_error_handling(setup_extractor):
    extractor, mock_loader = setup_extractor
    mock_loader.load_file.side_effect = Exception("File not found")
    with pytest.raises(Exception) as excinfo:
        extractor.load('error_pdf.pdf')
    assert "File not found" in str(excinfo.value)






import pytest
from unittest.mock import MagicMock

# Assuming FileExtractor and loader classes are already imported

@pytest.fixture
def mock_loader():
    """Fixture to provide a mock loader for DOCX files."""
    return MagicMock()

@pytest.fixture
def sample_docx(mock_loader):
    """Fixture to simulate a DOCX file with paragraphs and tables."""
    docx_mock = MagicMock()
    # Mocking the paragraphs in DOCX
    paragraph1 = MagicMock(text="This is the first paragraph.")
    paragraph2 = MagicMock(text="This is the second paragraph.")
    docx_mock.paragraphs = [paragraph1, paragraph2]
    
    # Mocking tables in DOCX
    cell1, cell2 = MagicMock(text="Cell1"), MagicMock(text="Cell2")
    row = MagicMock(cells=[cell1, cell2])
    table = MagicMock(rows=[row])
    docx_mock.tables = [table]
    
    # Setting the return value for loading file
    mock_loader.load_file.return_value = docx_mock
    return docx_mock

@pytest.fixture
def file_extractor(mock_loader):
    """Fixture to provide the FileExtractor instance for DOCX."""
    return FileExtractor(loader=mock_loader, file_type="docx")

### 10 Pytest Cases ###

def test_load_file(file_extractor, mock_loader):
    """Test if load function is called properly with file path."""
    file_extractor.load('test.docx')
    mock_loader.load_file.assert_called_with('test.docx')
    assert file_extractor.file_path == 'test.docx'

def test_extract_text_from_paragraphs(file_extractor, sample_docx):
    """Test extracting text from paragraphs in a DOCX file."""
    file_extractor.load('test.docx')
    text = file_extractor.extract_text()
    assert "This is the first paragraph." in text
    assert "This is the second paragraph." in text

def test_extract_text_from_table(file_extractor, sample_docx):
    """Test extracting text from a table in a DOCX file."""
    file_extractor.load('test.docx')
    text = file_extractor.extract_text()
    assert "Cell1\tCell2" in text

def test_empty_docx_file(file_extractor, mock_loader):
    """Test extracting text from an empty DOCX file."""
    empty_docx = MagicMock(paragraphs=[], tables=[])
    mock_loader.load_file.return_value = empty_docx
    file_extractor.load('empty.docx')
    text = file_extractor.extract_text()
    assert text == ""

def test_file_with_only_paragraphs(file_extractor, sample_docx):
    """Test DOCX file with only paragraphs and no tables."""
    sample_docx.tables = []
    file_extractor.load('paragraphs_only.docx')
    text = file_extractor.extract_text()
    assert "This is the first paragraph." in text
    assert "This is the second paragraph." in text

def test_file_with_only_tables(file_extractor, sample_docx):
    """Test DOCX file with only tables and no paragraphs."""
    sample_docx.paragraphs = []
    file_extractor.load('tables_only.docx')
    text = file_extractor.extract_text()
    assert "Cell1\tCell2" in text

def test_incorrect_file_type(mock_loader):
    """Test extracting text with incorrect file type raises error."""
    file_extractor = FileExtractor(loader=mock_loader, file_type="txt")
    file_extractor.load('test.txt')
    with pytest.raises(ValueError, match="Unsupported file type: txt"):
        file_extractor.extract_text()

def test_extract_text_with_multiple_paragraphs_and_tables(file_extractor, sample_docx):
    """Test DOCX file with multiple paragraphs and tables."""
    paragraph3 = MagicMock(text="This is another paragraph.")
    sample_docx.paragraphs.append(paragraph3)
    cell3, cell4 = MagicMock(text="Cell3"), MagicMock(text="Cell4")
    row2 = MagicMock(cells=[cell3, cell4])
    table2 = MagicMock(rows=[row2])
    sample_docx.tables.append(table2)

    file_extractor.load('complex.docx')
    text = file_extractor.extract_text()
    assert "This is another paragraph." in text
    assert "Cell3\tCell4" in text

def test_extract_text_with_special_characters(file_extractor, sample_docx):
    """Test DOCX file with special characters in paragraphs and tables."""
    sample_docx.paragraphs[0].text = "This is a paragraph with special char ©™."
    sample_docx.tables[0].rows[0].cells[0].text = "Cell ©"
    file_extractor.load('special_chars.docx')
    text = file_extractor.extract_text()
    assert "paragraph with special char ©™." in text
    assert "Cell ©" in text

def test_extract_text_with_empty_cells(file_extractor, sample_docx):
    """Test DOCX file with empty cells in tables."""
    sample_docx.tables[0].rows[0].cells[1].text = ""
    file_extractor.load('empty_cells.docx')
    text = file_extractor.extract_text()
    assert "Cell1\t" in text









import pytest
from unittest.mock import Mock

from pptx import Presentation

@pytest.fixture
def mock_loader():
    """Fixture to mock the file loader."""
    return Mock()

@pytest.fixture
def mock_ppt():
    """Fixture to create a mock pptx presentation."""
    ppt = Mock(spec=Presentation)
    return ppt

# Test case 1: Extract text from a single slide with text
def test_extract_text_single_slide(mock_loader, mock_ppt):
    slide = Mock()
    slide.shapes = [Mock(text="This is a slide text", has_table=False)]
    mock_ppt.slides = [slide]
    
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    extracted_text = extractor.extract_text()
    
    assert "This is a slide text" in extracted_text



# Test case 3: Extract text from multiple slides
def test_extract_text_multiple_slides(mock_loader, mock_ppt):
    slide1 = Mock()
    slide1.shapes = [Mock(text="Slide 1 text", has_table=False)]
    
    slide2 = Mock()
    slide2.shapes = [Mock(text="Slide 2 text", has_table=False)]
    
    mock_ppt.slides = [slide1, slide2]
    
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    extracted_text = extractor.extract_text()
    
    assert "Slide 1 text" in extracted_text
    assert "Slide 2 text" in extracted_text

# Test case 4: Extract images from slides
def test_extract_images(mock_loader, mock_ppt):
    slide = Mock()
    slide.shapes = [Mock(shape_type=13, image=Mock(blob=b'fake_image', ext="png"))]
    
    mock_ppt.slides = [slide]
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    images = extractor.extract_images()
    
    assert len(images) == 1
    assert images[0]['image_data'] == b'fake_image'
    assert images[0]['ext'] == "png"

# Test case 5: Extract URLs from slides with hyperlinks
def test_extract_urls_with_hyperlinks(mock_loader, mock_ppt):
    slide = Mock()
    shape = Mock()
    shape.text_frame = Mock(paragraphs=[Mock(runs=[Mock(text="Click here", hyperlink=Mock(address="https://example.com"))])])
    
    slide.shapes = [shape]
    mock_ppt.slides = [slide]
    
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    urls = extractor.extract_urls()
    
    assert len(urls) == 1
    assert urls[0]["linked_text"] == "Click here"
    assert urls[0]["url"] == "https://example.com"

# Test case 6: Extract URLs from slides without hyperlinks
def test_extract_urls_without_hyperlinks(mock_loader, mock_ppt):
    slide = Mock()
    shape = Mock()
    shape.text_frame = Mock(paragraphs=[Mock(runs=[Mock(text="No hyperlink", hyperlink=None)])])
    
    slide.shapes = [shape]
    mock_ppt.slides = [slide]
    
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    urls = extractor.extract_urls()
    
    assert len(urls) == 0




# Test case 8: Extract tables from slides without tables
def test_extract_tables_no_table(mock_loader, mock_ppt):
    shape = Mock()
    shape.has_table = False
    
    slide = Mock()
    slide.shapes = [shape]
    
    mock_ppt.slides = [slide]
    
    mock_loader.load_file.return_value = mock_ppt
    
    extractor = FileExtractor(mock_loader, file_type="pptx")
    extractor.load("dummy_path")
    
    tables = extractor.extract_tables()
    
    assert len(tables) == 0



# Test case 10: Handle unsupported file type
def test_unsupported_file_type(mock_loader):
    extractor = FileExtractor(mock_loader, file_type="txt")  # Unsupported file type
    
    with pytest.raises(ValueError, match="Unsupported file type"):
        extractor.extract_text()






# for sql storage test


@pytest.fixture
def mock_connection():
    """Create a mock SQLite connection."""
    return MagicMock()



@pytest.fixture
def sql_storage(mock_connection):
    """Create an instance of SQLStorage with a mocked connection."""
    with patch('sqlite3.connect', return_value=mock_connection):
        storage = SQLStorage(connection_string="mock_connection_string")
        yield storage




def test_store_inserts_data(sql_storage, mock_connection):
    """Test if the store method inserts data into the table."""
    sql_storage.store("TestTable", {"key": "value"})

    # Verify that the data was inserted into the table
    escaped_table_name = '"TestTable"'
    mock_connection.cursor().execute.assert_any_call(
        f"INSERT INTO {escaped_table_name} (data) VALUES (?)", (str({"key": "value"}),)
    )


def test_close_connection(sql_storage, mock_connection):
    """Test if the close method closes the database connection."""
    sql_storage.close()
    mock_connection.close.assert_called_once()



def test_store_creates_table_with_escaped_name(sql_storage, mock_connection):
    """Test if the store method creates a table with an escaped name."""
    sql_storage.store("Table Name With Space", {"key": "value"})

    # Check if the table was created with the correct escaped name
    escaped_table_name = '"Table_Name_With_Space"'
    
    # Assert that the CREATE TABLE statement was called
    mock_connection.cursor().execute.assert_any_call(
        f"""CREATE TABLE IF NOT EXISTS {escaped_table_name} (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data TEXT
        )"""
    )    



def test_store_inserts_data(sql_storage, mock_connection):
    """Test if the store method inserts data into the table."""
    sql_storage.store("TestTable", {"key": "value"})

    # Verify that the data was inserted into the table
    escaped_table_name = '"TestTable"'
    mock_connection.cursor().execute.assert_any_call(
        f"INSERT INTO {escaped_table_name} (data) VALUES (?)", (str({"key": "value"}),)
    )   



def test_close_connection(sql_storage, mock_connection):
    """Test if the close method closes the database connection."""
    sql_storage.close()
    mock_connection.close.assert_called_once()     







def test_store_inserts_empty_data(sql_storage, mock_connection):
    """Test if the store method can handle empty data."""
    sql_storage.store("EmptyTable", {})
    
    # Verify that the data was inserted into the table
    escaped_table_name = '"EmptyTable"'
    mock_connection.cursor().execute.assert_any_call(
        f"INSERT INTO {escaped_table_name} (data) VALUES (?)", (str({}),)
    )


def test_store_creates_table_with_special_characters(sql_storage, mock_connection):
    """Test if the store method creates a table with special characters in the name."""
    sql_storage.store("Table#With$Special%Characters", {"key": "value"})
    
    # Check if the table was created with the correct escaped name
    escaped_table_name = '"Table#With$Special%Characters"'
    mock_connection.cursor().execute.assert_any_call(
        f"""CREATE TABLE IF NOT EXISTS {escaped_table_name} (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data TEXT
        )"""
    )




def test_close_connection_is_called_after_operations(sql_storage, mock_connection):
    """Test if the close method is called after storing data."""
    sql_storage.store("CloseTestTable", {"key": "value"})
    sql_storage.close()
    
    # Assert that the close method was called
    mock_connection.close.assert_called_once()


def test_store_with_table_name_in_lowercase(sql_storage, mock_connection):
    """Test if the store method handles lowercase table names correctly."""
    sql_storage.store("lowercase_table", {"key": "value"})

    # Check if the table was created with the correct escaped name
    escaped_table_name = '"lowercase_table"'
    mock_connection.cursor().execute.assert_any_call(
        f"""CREATE TABLE IF NOT EXISTS {escaped_table_name} (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data TEXT
        )"""
    )


def test_store_inserts_large_data(sql_storage, mock_connection):
    """Test if the store method can handle large data."""
    large_data = "x" * 10000  # 10,000 characters
    sql_storage.store("LargeDataTable", {"large_key": large_data})

    escaped_table_name = '"LargeDataTable"'
    mock_connection.cursor().execute.assert_any_call(
        f"INSERT INTO {escaped_table_name} (data) VALUES (?)", (str({"large_key": large_data}),)
    )





